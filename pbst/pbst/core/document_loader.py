"""Utilities for loading source documents."""
import io
from pathlib import Path


def load_document(filename, data):

    extension = Path(
        filename
    ).suffix.lower()


    if extension == ".pdf":

        from pypdf import PdfReader

        reader = PdfReader(
            io.BytesIO(data)
        )

        text = []

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                text.append(
                    page_text
                )

        return "\n\n".join(text)


    elif extension == ".docx":

        from docx import Document

        document = Document(
            io.BytesIO(data)
        )

        text = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text.append(
                    paragraph.text
                )


        for table in document.tables:

            for row in table.rows:

                text.append(
                    " | ".join(
                        cell.text
                        for cell in row.cells
                    )
                )

        return "\n".join(text)


    elif extension in [
        ".txt",
        ".md"
    ]:

        return data.decode(
            "utf-8",
            errors="ignore"
        )


    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )